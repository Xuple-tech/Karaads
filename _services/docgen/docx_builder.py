import os
import re
import tempfile

from bs4 import NavigableString
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Inches, Pt, RGBColor

from markdown_utils import markdown_to_soup


def _slugify(value: str) -> str:
    slug = re.sub(r"[^a-zA-Z0-9]+", "-", value or "document").strip("-").lower()
    return slug or "document"


def _configure_styles(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(11)

    for style_name, size, color in [
        ("Heading 1", 22, RGBColor(15, 23, 42)),
        ("Heading 2", 16, RGBColor(15, 23, 42)),
        ("Heading 3", 13, RGBColor(29, 78, 216)),
    ]:
        style = document.styles[style_name]
        style.font.name = "Aptos"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color

    if "Kwati Code" not in document.styles:
        code_style = document.styles.add_style("Kwati Code", WD_STYLE_TYPE.PARAGRAPH)
        code_style.font.name = "Consolas"
        code_style.font.size = Pt(9)


def _append_inline(paragraph, node) -> None:
    for child in node.children:
        if isinstance(child, NavigableString):
            paragraph.add_run(str(child))
            continue

        text = child.get_text()
        run = paragraph.add_run(text)
        if child.name in ("strong", "b"):
            run.bold = True
        elif child.name in ("em", "i"):
            run.italic = True
        elif child.name == "code":
            run.font.name = "Consolas"
        elif child.name == "a":
            run.font.color.rgb = RGBColor(37, 99, 235)
            run.underline = True


def build_docx(title: str, content_markdown: str, options: dict):
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    _configure_styles(document)
    soup = markdown_to_soup(content_markdown)

    for node in soup.contents:
        name = getattr(node, "name", None)
        if name is None:
            continue
        if name == "h1":
            paragraph = document.add_paragraph(style="Heading 1")
            _append_inline(paragraph, node)
        elif name == "h2":
            paragraph = document.add_paragraph(style="Heading 2")
            _append_inline(paragraph, node)
        elif name == "h3":
            paragraph = document.add_paragraph(style="Heading 3")
            _append_inline(paragraph, node)
        elif name == "p":
            paragraph = document.add_paragraph()
            _append_inline(paragraph, node)
        elif name in ("ul", "ol"):
            style = "List Number" if name == "ol" else "List Bullet"
            for item in node.find_all("li", recursive=False):
                paragraph = document.add_paragraph(style=style)
                _append_inline(paragraph, item)
        elif name == "blockquote":
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.25)
            run = paragraph.add_run(node.get_text())
            run.italic = True
            run.font.color.rgb = RGBColor(71, 85, 105)
        elif name == "pre":
            paragraph = document.add_paragraph(style="Kwati Code")
            paragraph.add_run(node.get_text("\n").strip())
        elif name == "table":
            rows = node.find_all("tr")
            if not rows:
                continue
            col_count = max(len(row.find_all(["th", "td"])) for row in rows)
            table = document.add_table(rows=len(rows), cols=col_count)
            table.style = "Table Grid"
            for row_index, row in enumerate(rows):
                cells = row.find_all(["th", "td"])
                for col_index, cell in enumerate(cells):
                    table.cell(row_index, col_index).text = cell.get_text(" ", strip=True)

    if not soup.contents:
        document.add_heading(title, level=1)

    temp_dir = tempfile.mkdtemp(prefix="kwati_docgen_")
    filename = f"{_slugify(title)}.docx"
    output_path = os.path.join(temp_dir, filename)
    document.save(output_path)

    return output_path, filename
