#!/usr/bin/env python3
import sys
import json
import os
import re
import tempfile
import subprocess
import shutil
from pathlib import Path


def parse_markdown(text):
    lines = text.splitlines()
    blocks = []
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.strip().startswith('|'):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                table_lines.append(lines[i])
                i += 1
            rows = []
            for tline in table_lines:
                cells = [c.strip() for c in tline.strip().strip('|').split('|')]
                if all(re.match(r'^:?-+:?$', c) for c in cells if c):
                    continue
                rows.append(cells)
            if rows:
                blocks.append(('table', rows))
            continue
        if line.startswith('### '):
            blocks.append(('h3', line[4:].strip()))
        elif line.startswith('## '):
            blocks.append(('h2', line[3:].strip()))
        elif line.startswith('# '):
            blocks.append(('h1', line[2:].strip()))
        elif re.match(r'^[-*] ', line):
            blocks.append(('bullet', line[2:].strip()))
        elif re.match(r'^\d+\. ', line):
            blocks.append(('numbered', re.sub(r'^\d+\. ', '', line).strip()))
        elif line.strip() == '':
            blocks.append(('blank', ''))
        else:
            blocks.append(('para', line.strip()))
        i += 1
    return blocks


def strip_inline(text):
    text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
    text = re.sub(r'\*(.+?)\*', r'\1', text)
    text = re.sub(r'__(.+?)__', r'\1', text)
    text = re.sub(r'_(.+?)_', r'\1', text)
    text = re.sub(r'`(.+?)`', r'\1', text)
    return text


def build_docx(title, content_markdown, document_type, options):
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    formatting = options.get('formatting', {}) or {}
    font_family = formatting.get('font_family', 'Inter')
    font_size = int(formatting.get('font_size', 14) or 14)
    text_color = str(formatting.get('text_color', '#1F2937')).lstrip('#')
    paper_size = str(formatting.get('paper_size', 'a4')).lower()
    if not re.match(r'^[0-9A-Fa-f]{6}$', text_color):
        text_color = '1F2937'
    rgb_color = RGBColor.from_string(text_color.upper())
    page_sizes = {
        'a4': (8.27, 11.69),
        'letter': (8.5, 11.0),
        'legal': (8.5, 14.0),
    }
    page_width, page_height = page_sizes.get(paper_size, page_sizes['a4'])

    for section in doc.sections:
        section.page_width = Inches(page_width)
        section.page_height = Inches(page_height)
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1.2)
        section.right_margin = Inches(1.2)

    title_para = doc.add_heading(title, level=0)
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in title_para.runs:
        run.font.name = font_family
        run.font.size = Pt(max(font_size + 10, 20))
        run.font.color.rgb = rgb_color

    if options.get('include_header', True):
        from datetime import date
        date_para = doc.add_paragraph(date.today().strftime('%B %d, %Y'))
        date_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        date_para.runs[0].font.size = Pt(10)
        date_para.runs[0].font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    doc.add_paragraph()

    blocks = parse_markdown(content_markdown)

    for style, data in blocks:
        if style == 'blank':
            continue
        elif style == 'table':
            rows = data
            if not rows:
                continue
            num_cols = max(len(r) for r in rows)
            tbl = doc.add_table(rows=len(rows), cols=num_cols)
            tbl.style = 'Table Grid'
            for r_idx, row_data in enumerate(rows):
                row = tbl.rows[r_idx]
                for c_idx, cell_text in enumerate(row_data):
                    if c_idx < num_cols:
                        cell = row.cells[c_idx]
                        cell.text = strip_inline(cell_text)
                        if r_idx == 0:
                            for run in cell.paragraphs[0].runs:
                                run.bold = True
            doc.add_paragraph()
        else:
            text = strip_inline(data)
            if style == 'h1':
                para = doc.add_heading(text, level=1)
            elif style == 'h2':
                para = doc.add_heading(text, level=2)
            elif style == 'h3':
                para = doc.add_heading(text, level=3)
            elif style == 'bullet':
                para = doc.add_paragraph(text, style='List Bullet')
            elif style == 'numbered':
                para = doc.add_paragraph(text, style='List Number')
            else:
                para = doc.add_paragraph(text)

            for run in para.runs:
                run.font.name = font_family
                if style == 'h1':
                    run.font.size = Pt(max(font_size + 6, 18))
                elif style == 'h2':
                    run.font.size = Pt(max(font_size + 4, 16))
                elif style == 'h3':
                    run.font.size = Pt(max(font_size + 2, 15))
                else:
                    run.font.size = Pt(font_size)
                run.font.color.rgb = rgb_color

    tmp = tempfile.NamedTemporaryFile(suffix='.docx', delete=False)
    tmp.close()
    doc.save(tmp.name)
    return tmp.name


def convert_docx_to_pdf(docx_path):
    """Convert a DOCX file to PDF using LibreOffice headless."""
    out_dir = tempfile.mkdtemp()
    try:
        result = subprocess.run(
            [
                'libreoffice', '--headless', '--convert-to', 'pdf',
                '--outdir', out_dir, docx_path,
            ],
            capture_output=True,
            text=True,
            timeout=60,
        )
        if result.returncode != 0:
            raise RuntimeError(f'LibreOffice conversion failed: {result.stderr}')

        pdf_files = list(Path(out_dir).glob('*.pdf'))
        if not pdf_files:
            raise RuntimeError('LibreOffice did not produce a PDF file.')

        # Move to a stable temp path
        out_pdf = tempfile.NamedTemporaryFile(suffix='.pdf', delete=False)
        out_pdf.close()
        shutil.move(str(pdf_files[0]), out_pdf.name)
        return out_pdf.name
    finally:
        shutil.rmtree(out_dir, ignore_errors=True)


def generate_docx(title, content_markdown, document_type, options):
    return build_docx(title, content_markdown, document_type, options)


def generate_pdf(title, content_markdown, document_type, options):
    docx_path = build_docx(title, content_markdown, document_type, options)
    try:
        return convert_docx_to_pdf(docx_path)
    finally:
        os.unlink(docx_path)


def main():
    raw = sys.stdin.read().strip()
    try:
        payload = json.loads(raw)
    except Exception as e:
        print(json.dumps({'success': False, 'error': f'Invalid JSON input: {e}'}))
        sys.exit(1)

    title = payload.get('title', 'Document')
    content = payload.get('content_markdown', '')
    fmt = payload.get('format', 'docx').lower()
    doc_type = payload.get('document_type', 'general')
    options = payload.get('options', {})

    try:
        if fmt == 'docx':
            output_path = generate_docx(title, content, doc_type, options)
            mime = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        elif fmt == 'pdf':
            output_path = generate_pdf(title, content, doc_type, options)
            mime = 'application/pdf'
        else:
            raise ValueError(f'Unsupported format: {fmt}')

        slug = re.sub(r'[^a-z0-9]+', '-', title.lower()).strip('-')[:40]
        filename = f'{slug}.{fmt}'

        print(json.dumps({
            'success': True,
            'output_path': output_path,
            'filename': filename,
            'mime_type': mime,
        }))
    except Exception as e:
        print(json.dumps({'success': False, 'error': str(e)}))
        sys.exit(1)


if __name__ == '__main__':
    main()
